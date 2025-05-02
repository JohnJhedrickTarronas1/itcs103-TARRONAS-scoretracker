import os
import openpyxl
from openpyxl import Workbook, load_workbook
from tkinter import Tk, Frame, Label, Entry, Button, messagebox, ttk

FILENAME = "student_scores.xlsx"

# workbook
def init_workbook():
    if not os.path.exists(FILENAME):
        wb = Workbook()
        ws = wb.active
        ws.title = "Scores"
        ws.append(["Name", "Score 1", "Score 2", "Score 3", "Average", "Status"])
        wb.save(FILENAME)

# Determine status
def get_status(avg):
    return "Pass" if avg >= 75 else "Fail"

# Add or update student
def add_or_update_student(name, s1, s2, s3):
    try:
        s1 = float(s1)
        s2 = float(s2)
        s3 = float(s3)
    except ValueError:
        messagebox.showerror("Invalid Input", "All scores must be numbers.")
        return

    average = round((s1 + s2 + s3) / 3, 2)
    status = get_status(average)

    wb = load_workbook(FILENAME)
    ws = wb.active

    for row in ws.iter_rows(min_row=2):
        if row[0].value == name:
            row[1].value = s1
            row[2].value = s2
            row[3].value = s3
            row[4].value = average
            row[5].value = status
            wb.save(FILENAME)
            messagebox.showinfo("Updated", f"{name}'s record updated.")
            return

    ws.append([name, s1, s2, s3, average, status])
    wb.save(FILENAME)
    messagebox.showinfo("Added", f"{name} added successfully.")

# Display records
def display_all_records():
    wb = load_workbook(FILENAME)
    ws = wb.active
    records_list.delete(*records_list.get_children())
    for row in ws.iter_rows(min_row=2, values_only=True):
        records_list.insert("", "end", values=row)

# Submit button 
def submit_record():
    name = name_entry.get()
    s1 = score1_entry.get()
    s2 = score2_entry.get()
    s3 = score3_entry.get()
    if not name or not s1 or not s2 or not s3:
        messagebox.showwarning("Missing Data", "All fields are required.")
        return
    add_or_update_student(name, s1, s2, s3)
    name_entry.delete(0, 'end')
    score1_entry.delete(0, 'end')
    score2_entry.delete(0, 'end')
    score3_entry.delete(0, 'end')
    display_all_records()

# Delete student record
def delete_selected_record():
    selected = records_list.selection()
    if not selected:
        messagebox.showwarning("No selection", "Please select a record to delete.")
        return

    name = records_list.item(selected[0])['values'][0]
    wb = load_workbook(FILENAME)
    ws = wb.active

    for row in ws.iter_rows(min_row=2):
        if row[0].value == name:
            ws.delete_rows(row[0].row, 1)
            wb.save(FILENAME)
            messagebox.showinfo("Deleted", f"{name}'s record deleted.")
            display_all_records()
            return

    messagebox.showerror("Not Found", f"{name} not found.")

# GUI
init_workbook()
windows = Tk()
windows.title("Student Score Tracker")
windows.geometry("500x400")

# Input frame
input_frame = Frame(windows)
input_frame.pack(pady=10)

Label(input_frame, text="Name:").grid(row=0, column=0, padx=5, pady=5, sticky="e")
name_entry = Entry(input_frame, width=20)
name_entry.grid(row=0, column=1, padx=5, pady=5)

Label(input_frame, text="Score 1:").grid(row=1, column=0, padx=5, pady=5, sticky="e")
score1_entry = Entry(input_frame, width=20)
score1_entry.grid(row=1, column=1, padx=5, pady=5)

Label(input_frame, text="Score 2:").grid(row=2, column=0, padx=5, pady=5, sticky="e")
score2_entry = Entry(input_frame, width=20)
score2_entry.grid(row=2, column=1, padx=5, pady=5)

Label(input_frame, text="Score 3:").grid(row=3, column=0, padx=5, pady=5, sticky="e")
score3_entry = Entry(input_frame, width=20)
score3_entry.grid(row=3, column=1, padx=5, pady=5)

Button(windows, text="Add/Update", command=submit_record).pack(pady=5)
Button(windows, text="Delete Selected", command=delete_selected_record).pack(pady=5)

# Treeview for displaying records
records_frame = Frame(windows)
records_frame.pack(fill="both", expand=True, padx=10)

columns = ("Name", "Score 1", "Score 2", "Score 3", "Average", "Status")
records_list = ttk.Treeview(records_frame, columns=columns, show="headings")
for col in columns:
    records_list.heading(col, text=col)
    records_list.column(col, anchor="center")

records_list.pack(fill="both", expand=True)

display_all_records()
windows.mainloop()
